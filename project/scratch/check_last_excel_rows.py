import openpyxl

wb = openpyxl.load_workbook(r"f:\BILLING\Bimonthly_Bills_2021_to_August_2026.xlsx", data_only=True)
sheet = wb.worksheets[0]
for r in range(max(2, sheet.max_row - 10), sheet.max_row + 1):
    row_vals = [cell.value for cell in sheet[r]]
    # Column 0: Customer_ID, 1: Consumer_Name, 8: Billing_Month, 17: Units
    print(f"Row {r}: ID={row_vals[0]}, Month={row_vals[8]}, Units={row_vals[17]}")
