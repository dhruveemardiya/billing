import openpyxl

wb = openpyxl.load_workbook(r"f:\BILLING\Bimonthly_Bills_2021_to_August_2026.xlsx", data_only=True)
sheet = wb.worksheets[0]
print("Sheet title:", sheet.title)
headers = [cell.value for cell in sheet[1]]
print("Headers:", headers)
print(f"Total rows: {sheet.max_row}")
for r in range(2, min(sheet.max_row + 1, 15)):
    row_vals = [cell.value for cell in sheet[r]]
    print(f"Row {r}: customer_id={row_vals[1] if len(row_vals)>1 else None}, month={row_vals[6] if len(row_vals)>6 else None}, units={row_vals[10] if len(row_vals)>10 else None}, row={row_vals[:8]}")
